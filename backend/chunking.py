"""Split code into overlapping line windows and normalise for embedding."""
import re

CODE_EXT = {".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".java", ".rb", ".rs", ".c", ".h",
            ".cpp", ".cs", ".php", ".kt", ".swift", ".scala", ".sh", ".sql"}
SKIP_DIRS = {".git", "node_modules", "venv", ".venv", "__pycache__", "dist", "build", ".tox"}


_ZW = re.compile("[\u200b-\u200f\u2060-\u2064\ufeff\u00ad\u034f\u180e]")
# Latin look-alikes that NFKC leaves alone (Cyrillic and Greek). Small on purpose: the common swaps only.
_HOMO = str.maketrans("\u0430\u0435\u043e\u0440\u0441\u0443\u0445\u0456\u0458\u0410\u0412\u0415\u041a\u041c\u041d\u041e\u0420\u0421\u0422\u0425\u03bf\u039f\u0391\u0392\u0395\u0397\u0399\u039a\u039c\u039d\u03a1\u03a4\u03a7",
                      "aeopcyxijABEKMHOPCTXoOABEHIKMNPTX")


def clean(text: str) -> str:
    """Undo cheap evasion tricks before matching: zero-width characters, full-width and other
    compatibility forms (NFKC), and the common Cyrillic or Greek look-alikes of Latin letters."""
    import unicodedata
    return unicodedata.normalize("NFKC", _ZW.sub("", text)).translate(_HOMO)


def normalise(text: str) -> str:
    # Collapse whitespace so reformatting does not change the vector much.
    return re.sub(r"\s+", " ", clean(text)).strip()


def windows(text: str, size: int = 30, step: int = 15, min_chars: int = 40):
    """Yield (start_line, end_line, chunk) windows, 1-based inclusive."""
    lines = text.splitlines()
    if not lines:
        return
    i = 0
    while True:
        part = lines[i:i + size]
        chunk = "\n".join(part)
        if len(chunk.strip()) >= min_chars:
            yield i + 1, i + len(part), chunk
        if i + size >= len(lines):
            break
        i += step


DOC_EXT = {".md", ".txt", ".pdf", ".docx"}
# Plain text types the extension can also send to /check_file (the drop and picker paths).
TEXT_EXT = {".md", ".markdown", ".txt", ".csv", ".json", ".log", ".yaml", ".yml", ".html"} | CODE_EXT
UPLOAD_EXT = DOC_EXT | TEXT_EXT


class NotText(ValueError):
    pass


def doc_chunks(text: str, target: int = 160, max_words: int = 200, min_words: int = 120, overlap: int = 30):
    """Paragraph-aware chunks of about 120-200 words with ~30 words overlap.

    Yields (start_para, end_para, chunk). Paragraphs are packed until the
    target is reached; oversized paragraphs are split on word boundaries.
    """
    paras = [re.sub(r"[ \t]+", " ", p).strip() for p in re.split(r"\n\s*\n", text)]
    paras = [p for p in paras if p]
    units = []  # (para_index, words)
    for i, p in enumerate(paras):
        w = p.split()
        while len(w) > max_words:  # packing below adds the overlap, so split cleanly here
            units.append((i, w[:target]))
            w = w[target:]
        units.append((i, w))
    buf, start, last = [], None, None
    for i, w in units:
        if buf and len(buf) + len(w) > max_words and len(buf) >= min_words // 2:
            yield start + 1, last + 1, " ".join(buf)
            tail = buf[-overlap:]
            buf, start = list(tail), last
        if start is None:
            start = i
        buf.extend(w)
        last = i
        if len(buf) >= target:
            yield start + 1, last + 1, " ".join(buf)
            buf, start = buf[-overlap:], i
    if buf and (len(buf) > overlap or start == 0):
        yield (start or 0) + 1, (last or 0) + 1, " ".join(buf)


def _user_y(cm, tm) -> float:
    """Text origin in page space. cm and tm are the six PDF matrix numbers."""
    return cm[1] * tm[4] + cm[3] * tm[5] + cm[5]


def _paragraph_breaks(pieces) -> str | None:
    """Put back blank lines where the page leaves a gap bigger than one line.

    pypdf's plain extractor keeps one newline per visual line and drops the extra
    space between paragraphs, so a code block in a longer PDF is glued to the
    prose around it. The words stay in the same order; only newlines are added.
    Returns None when the positions are not usable.
    """
    flat = []
    for y, text in pieces:
        bits = text.split("\n")
        for i, bit in enumerate(bits):
            if bit:
                flat.append((y, bit))
            if i < len(bits) - 1:
                flat.append((None, "\n"))
    content = [(i, y) for i, (y, t) in enumerate(flat) if t.strip() and y is not None]
    gaps = []
    for (i1, y1), (i2, y2) in zip(content, content[1:]):
        gap = y1 - y2
        if gap > 2:
            gaps.append((i1, gap))
    if len(gaps) < 3:
        return None
    ordered = sorted(g for _, g in gaps)
    tight = ordered[: max(1, (len(ordered) + 1) // 2)]
    line_h = tight[len(tight) // 2]
    if line_h < 4:
        return None
    cut = max(line_h * 1.7, line_h + 6)
    break_after = {i for i, gap in gaps if gap >= cut}
    out, prev, used = [], None, set()
    for i, (y, t) in enumerate(flat):
        if t.strip():
            prev = i
            out.append(t)
        elif t == "\n" and prev in break_after and prev not in used:
            used.add(prev)
            out.append("\n\n")
        else:
            out.append(t)
    return "".join(out)


def _pdf_text(data: bytes, keep_paragraphs: bool) -> str:
    import io
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")  # owner-password-only PDFs open with an empty user password
        except Exception:
            ok = 0
        if not ok:
            raise NotText("the PDF is password protected")
    pages = []
    for page in reader.pages:
        if not keep_paragraphs:
            pages.append(page.extract_text() or "")
            continue
        pieces = []

        def visit(text, cm, tm, font, size, pieces=pieces):
            pieces.append((_user_y(cm, tm), text))

        plain = page.extract_text(visitor_text=visit) or ""
        restored = _paragraph_breaks(pieces)
        # Only newlines may change. If the positions produce different words, keep the plain extract.
        if restored is None or re.sub(r"\s+", "", restored) != re.sub(r"\s+", "", plain):
            pages.append(plain)
        else:
            pages.append(restored)
    return "\n\n".join(pages)


def extract_text(path_or_bytes, suffix: str, keep_paragraphs: bool = False) -> str:
    """Extract plain text from .md/.txt/.pdf/.docx (path or raw bytes).

    keep_paragraphs: for a PDF the user dropped, put back blank lines where the
    page has a gap. Off for the index, which was tuned on the plain extract.
    """
    suffix = suffix.lower()
    if isinstance(path_or_bytes, (bytes, bytearray)):
        data = bytes(path_or_bytes)
    else:
        with open(path_or_bytes, "rb") as f:
            data = f.read()
    if suffix == ".pdf":
        return _pdf_text(data, keep_paragraphs)
    if suffix == ".docx":
        import io
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n\n".join(parts)
    if b"\x00" in data[:8192]:
        raise NotText("the file looks binary, not text")
    return data.decode("utf-8", errors="replace").replace("\ufffd", "")


def doc_passages(text: str):
    """Two granularities for documents: ~120-200 word chunks (context) plus
    ~60 word passages (so a short excerpt or a one-paragraph summary still has
    a close neighbour). Yields (start_para, end_para, chunk)."""
    yield from doc_chunks(text)
    yield from doc_chunks(text, target=60, max_words=80, min_words=40, overlap=20)


# Paste-side windows (used by the backend and by the eval and benchmark scripts).
def prose_spans(text: str, words: int = 80, step: int = 40):
    """Character spans of paste prose: paragraphs, long ones split into ~80 word windows."""
    spans = []
    for m in re.finditer(r"\S(?:.*?\S)?(?=\n\s*\n|\s*$)", text, flags=re.S):
        toks = list(re.finditer(r"\S+", m.group(0)))
        if len(toks) < 6:
            continue
        base = m.start()
        if len(toks) <= words:
            spans.append((m.start(), m.end()))
            continue
        i = 0
        while True:
            part = toks[i:i + words]
            spans.append((base + part[0].start(), base + part[-1].end()))
            if i + words >= len(toks):
                break
            i += step
    return spans


def line_spans(text: str):
    """Character spans of 30-line windows (code mode)."""
    offs = [0]
    for ln in text.splitlines(keepends=True):
        offs.append(offs[-1] + len(ln))
    out = []
    for s, e, _ in windows(text, size=30, step=15, min_chars=40):
        out.append((offs[s - 1], offs[e]))
    return out


_TOK = re.compile(r"\w+")


def shingles(text: str, n: int = 3) -> set:
    """Hashed word n-gram fingerprint (lower-cased word 3-grams), used for exact and near-exact reuse."""
    import hashlib
    w = [x.lower() for x in _TOK.findall(clean(text))]
    return {int.from_bytes(hashlib.blake2b(" ".join(w[i:i + n]).encode(), digest_size=8).digest(), "big")
            for i in range(len(w) - n + 1)}

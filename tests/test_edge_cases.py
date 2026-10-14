"""Edge cases for /check and /check_file against the real model and index (no mocks)."""
import io
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from fastapi.testclient import TestClient  # noqa: E402

import app  # noqa: E402
from chunking import clean  # noqa: E402

client = TestClient(app.app)
ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "demo" / "company-docs"
MEMO = (DOCS / "board-memo-acquisition-ferrolake.md").read_text()
PUBLIC = ("The mitochondria is the powerhouse of the cell. Cellular respiration converts glucose and oxygen "
          "into carbon dioxide, water and ATP, which the cell uses as an energy source for most processes.")


def check(text, **kw):
    r = client.post("/check", json={"text": text, **kw})
    assert r.status_code == 200, r.text
    return r.json()


def upload(name, data):
    return client.post("/check_file", files={"file": (name, data)}, data={"source": "claude.ai"})


def memo_paragraph():
    return max(MEMO.split("\n\n"), key=len)


# ---------- text pastes ----------
def test_empty_and_whitespace_allow():
    for t in ("", "   ", "\n\n\t\n", "\u200b\u200b\u200b" * 20):
        r = check(t)
        assert r["verdict"] == "allow" and r["matches"] == []


def test_missing_text_field_is_422():
    assert client.post("/check", json={}).status_code == 422
    assert client.post("/check", content=b"not json", headers={"content-type": "application/json"}).status_code == 422


def test_over_limit_paste_rejected():
    assert client.post("/check", json={"text": "a" * (app.MAX_CHARS + 1)}).status_code == 422


def test_huge_paste_is_bounded_and_sees_the_end():
    filler = "\n\n".join(PUBLIC + f" Note {i}." for i in range(5000))  # about 1 MB of unrelated prose
    text = filler + "\n\n" + memo_paragraph()
    assert len(text) > 900_000
    t0 = time.time()
    r = check(text)
    assert time.time() - t0 < 30
    assert r["verdict"] in {"coach", "block"}, "confidential paragraph at the end of a long paste was missed"


def test_binary_like_text_does_not_crash():
    r = check("".join(chr(i) for i in range(1, 32)) * 10 + "\x00\x01\xff garbage \udcff".encode("utf-8", "surrogatepass").decode("utf-8", "replace"))
    assert r["verdict"] in {"allow", "nudge", "coach"}


def test_non_english_and_emoji_allow():
    for t in ("Bonjour, pouvez-vous m'aider à écrire une lettre de motivation pour un poste d'ingénieur logiciel ? Merci beaucoup.",
              "今日はとても良い天気ですね。週末に公園で友達とピクニックをする予定です。お弁当を作ります。",
              "Привет! Помоги мне, пожалуйста, составить список покупок на неделю для семьи из четырёх человек.",
              "मुझे एक अच्छी किताब के बारे में बताइए जो मैं इस सप्ताहांत पढ़ सकूं, कृपया कुछ सुझाव दें।",
              "🎉🎉 Party planning: balloons 🎈, cake 🎂, music 🎶 and games 🎲 for my daughter's 7th birthday!! 🥳"):
        assert check(t)["verdict"] == "allow", t


def test_mixed_language_with_confidential_doc_still_caught():
    text = "Résume ce document s'il te plaît / 请总结一下 🙏\n\n" + memo_paragraph()
    assert check(text)["verdict"] in {"coach", "block"}


def test_zero_width_and_homoglyph_evasion_caught():
    para = memo_paragraph()
    zw = "\u200b".join(para)  # a zero-width space between every character
    homo = para.replace("a", "\u0430").replace("o", "\u043e").replace("e", "\u0435")  # Cyrillic look-alikes
    base = check(para)["verdict"]
    assert base in {"coach", "block"}
    assert check(zw)["verdict"] == base
    assert check(homo)["verdict"] in {"coach", "block"}


def test_clean_is_cheap_and_idempotent():
    s = "Ｃａｌｌｏｗａｙ\u200b Metrics \u0441\u043eding"
    assert clean(s) == "Calloway Metrics coding" and clean(clean(s)) == clean(s)


def test_unrelated_code_in_several_languages_allows():
    samples = {
        "js": "function debounce(fn, ms) {\n  let t;\n  return (...args) => {\n    clearTimeout(t);\n    t = setTimeout(() => fn(...args), ms);\n  };\n}",
        "go": "package main\n\nimport \"fmt\"\n\nfunc fib(n int) int {\n\tif n < 2 {\n\t\treturn n\n\t}\n\treturn fib(n-1) + fib(n-2)\n}\n\nfunc main() { fmt.Println(fib(10)) }",
        "rust": "fn gcd(mut a: u64, mut b: u64) -> u64 {\n    while b != 0 {\n        let t = b;\n        b = a % b;\n        a = t;\n    }\n    a\n}",
        "sql": "SELECT customer_id, COUNT(*) AS orders\nFROM orders\nWHERE created_at >= DATE '2024-01-01'\nGROUP BY customer_id\nHAVING COUNT(*) > 5\nORDER BY orders DESC;",
    }
    for lang, code in samples.items():
        assert check(code)["verdict"] in {"allow", "nudge"}, lang


def test_secret_with_unicode_around_it_blocks():
    r = check("🔑 clé: AKIAABCDEFGHIJKLMNOP — ne pas partager")
    assert r["verdict"] == "block" and "AKIA" not in r["redacted"]


def test_concurrent_requests():
    texts = [PUBLIC, memo_paragraph(), "AKIAABCDEFGHIJKLMNOP is the key"] * 8
    with ThreadPoolExecutor(8) as ex:
        out = list(ex.map(lambda t: client.post("/check", json={"text": t}).json()["verdict"], texts))
    assert out[0::3] == ["allow"] * 8
    assert set(out[1::3]) <= {"coach", "block"} and len(set(out[1::3])) == 1
    assert out[2::3] == ["block"] * 8


# ---------- files (picker and drag-and-drop both post to /check_file) ----------
def test_drop_text_files_supported():
    for name in ("notes.md", "export.csv", "snippet.py", "config.json"):
        r = upload(name, memo_paragraph().encode() if name == "notes.md" else PUBLIC.encode())
        assert r.status_code == 200, (name, r.text)
    assert upload("notes.md", memo_paragraph().encode()).json()["verdict"] in {"coach", "block"}


def test_unsupported_types_415():
    for name in ("photo.png", "clip.mp4", "archive.zip", "noext", "setup.exe"):
        assert upload(name, b"\x89PNG\r\n\x1a\n" + b"\x00" * 64).status_code == 415, name


def test_binary_disguised_as_text_422():
    r = upload("secret.txt", b"\x89PNG\r\n\x1a\n\x00\x00\x00IHDR" + bytes(range(256)))
    assert r.status_code == 422 and "binary" in r.json()["detail"]


def test_empty_text_file_allows():
    r = upload("empty.txt", b"")
    assert r.status_code == 200 and r.json()["verdict"] == "allow" and r.json()["file"]["chars"] == 0


def test_corrupt_pdf_and_docx_422():
    assert upload("broken.pdf", b"%PDF-1.7\nthis is not really a pdf").status_code == 422
    assert upload("broken.docx", b"PK\x03\x04 not a zip").status_code == 422
    assert upload("empty.pdf", b"").status_code == 422


def test_encrypted_pdf_422_with_clear_reason():
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for p in PdfReader(DOCS / "layoff-plan-q4-2026.pdf").pages:
        w.add_page(p)
    w.encrypt(user_password="hunter2", owner_password="owner")
    buf = io.BytesIO(); w.write(buf)
    r = upload("locked.pdf", buf.getvalue())
    assert r.status_code == 422 and "password" in r.json()["detail"]


def test_owner_only_encrypted_pdf_is_still_checked():
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for p in PdfReader(DOCS / "layoff-plan-q4-2026.pdf").pages:
        w.add_page(p)
    w.encrypt(user_password="", owner_password="owner")
    buf = io.BytesIO(); w.write(buf)
    r = upload("restricted.pdf", buf.getvalue())
    assert r.status_code == 200 and r.json()["verdict"] == "block"


def test_large_pdf_checked_quickly():
    from pypdf import PdfReader, PdfWriter
    src = PdfReader(DOCS / "layoff-plan-q4-2026.pdf")
    w = PdfWriter()
    for _ in range(60):
        for p in src.pages:
            w.add_page(p)
    buf = io.BytesIO(); w.write(buf)
    t0 = time.time()
    r = upload("big-layoff-plan.pdf", buf.getvalue())
    assert r.status_code == 200 and r.json()["verdict"] == "block"
    assert time.time() - t0 < 30


def test_large_docx_with_confidential_section_at_end():
    import docx
    d = docx.Document()
    for i in range(1500):
        d.add_paragraph(PUBLIC + f" Section {i}.")
    d.add_paragraph(memo_paragraph())
    buf = io.BytesIO(); d.save(buf)
    r = upload("handbook.docx", buf.getvalue())
    assert r.status_code == 200 and r.json()["verdict"] in {"coach", "block"}


def test_file_over_size_limit_413(monkeypatch):
    monkeypatch.setattr(app, "MAX_FILE_BYTES", 1000)
    assert upload("big.txt", b"a " * 1000).status_code == 413


def test_file_log_has_site_and_no_content():
    upload("layoff-plan-q4-2026.pdf", (DOCS / "layoff-plan-q4-2026.pdf").read_bytes())
    row = client.get("/admin/api").json()["recent"][0]
    assert row["site"] == "claude.ai" and row["channel"] == "file"
    assert len(row["sha256"]) == 64 and "text" not in row and "redacted" not in row


def test_health_reports_accurate_default():
    h = client.get("/health").json()
    assert h["profile"] == "accurate" and h["tfidf"]["enabled"] is True
    assert h["version"] == (ROOT / "VERSION").read_text().strip()
